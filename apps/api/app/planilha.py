"""Leitura de planilhas para prospecção em lote (CONSTITUICAO §7.2, entrada 2).

- `.csv` (UTF-8, com ou sem BOM, ou Latin-1; separador `,`, `;` ou tabulação) e `.xlsx` (só a 1ª aba, valores
  e não fórmulas). `.xlsm` e outros formatos são recusados.
- Limites: 5 MB, 1.000 linhas, 50 colunas e, no `.xlsx`, 50 MB descompactados (proteção contra "zip bomb").
- O mapeamento de colunas começa por regras sobre o nome do cabeçalho; o que faltar, o worker completa com IA.
- Cada linha vira nova, já cadastrada, repetida na planilha ou inválida, com o motivo em português.
"""

import csv
import io
import re
import unicodedata
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Literal

from app.domain import EntradaInvalida, normalizar_dominio, normalizar_nome

MAX_BYTES = 5 * 1024 * 1024
MAX_DESCOMPACTADO = 50 * 1024 * 1024
MAX_LINHAS = 1000
MAX_COLUNAS = 50
MAX_CELULA = 300

CAMPOS = ("empresa", "site", "cnpj", "cidade", "uf", "contato_nome", "contato_cargo", "contato_email")
ROTULO_CAMPO = {
    "empresa": "Empresa",
    "site": "Site",
    "cnpj": "CNPJ",
    "cidade": "Cidade",
    "uf": "UF",
    "contato_nome": "Nome do contato",
    "contato_cargo": "Cargo do contato",
    "contato_email": "E-mail do contato",
}
UFS = frozenset("AC AL AP AM BA CE DF ES GO MA MT MS MG PA PB PR PE PI RJ RN RS RO RR SC SP SE TO".split())
EMAIL = re.compile(r"^[^@\s]+@([a-z0-9-]+\.)+[a-z]{2,}$")

# Regras por nome de cabeçalho (sem acento, minúsculo). A ordem define a prioridade de cada campo.
PADROES: dict[str, list[str]] = {
    "cnpj": [r"\bcnpj\b"],
    "contato_email": [r"e-?mail", r"\bcorreio\b"],
    "contato_cargo": [r"\bcargo\b", r"\bfuncao\b", r"\btitle\b", r"\bposicao\b"],
    "contato_nome": [r"\bcontato\b", r"\bresponsavel\b", r"\bdecisor\b", r"nome do contato"],
    "site": [r"\bsite\b", r"\bwebsite\b", r"\burl\b", r"\bdominio\b", r"\bhomepage\b", r"\bendereco web\b"],
    "empresa": [
        r"razao social",
        r"nome fantasia",
        r"\bempresa\b",
        r"\borganizacao\b",
        r"\binstituicao\b",
        r"\bcliente\b",
        r"\bcompany\b",
        r"^nome$",
    ],
    "uf": [r"^uf$", r"^estado$", r"^state$"],
    "cidade": [r"\bcidade\b", r"\bmunicipio\b", r"\bcity\b"],
}


class PlanilhaInvalida(ValueError):
    """Mensagem em português, pronta para mostrar ao usuário."""


@dataclass
class Planilha:
    cabecalhos: list[str]
    linhas: list[list[str]]  # linhas em branco no meio ficam, para a numeração bater com a do arquivo
    linha_cabecalho: int = 1  # número da linha do cabeçalho no arquivo (a 1ª não vazia)

    @property
    def total(self) -> int:
        return sum(1 for linha in self.linhas if any(linha))


def _texto_celula(valor: object) -> str:
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)  # CNPJ e números lidos como 1.2345678000199e+13
    if isinstance(valor, datetime | date):
        return valor.isoformat()[:10]
    return " ".join(str(valor).split())[:MAX_CELULA]


def _ler_xlsx(dados: bytes) -> list[list[str]]:
    from openpyxl import load_workbook

    try:
        with zipfile.ZipFile(io.BytesIO(dados)) as z:
            if sum(i.file_size for i in z.infolist()) > MAX_DESCOMPACTADO:
                raise PlanilhaInvalida("A planilha é grande demais depois de descompactada.")
            if any(i.filename.lower().endswith("vbaproject.bin") for i in z.infolist()):
                raise PlanilhaInvalida("Planilhas com macros não são aceitas. Salve como .xlsx ou .csv.")
        livro = load_workbook(io.BytesIO(dados), read_only=True, data_only=True)
    except PlanilhaInvalida:
        raise
    except Exception as exc:  # noqa: BLE001 — arquivo corrompido ou que não é xlsx
        raise PlanilhaInvalida("Não foi possível ler o arquivo .xlsx. Confira se ele abre no Excel.") from exc
    try:
        aba = livro.worksheets[0]
        linhas = []
        for linha in aba.iter_rows(values_only=True):
            linhas.append([_texto_celula(v) for v in linha[:MAX_COLUNAS]])
            if len(linhas) > MAX_LINHAS + 200:
                break
        return linhas
    finally:
        livro.close()


def _ler_csv(dados: bytes) -> list[list[str]]:
    for codificacao in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            texto = dados.decode(codificacao)
            break
        except UnicodeDecodeError:
            continue
    if "\x00" in texto:
        raise PlanilhaInvalida("O arquivo .csv não parece ser texto. Exporte a planilha como CSV de novo.")
    amostra = texto[:4096]
    try:
        dialeto = csv.Sniffer().sniff(amostra, delimiters=",;\t")
        separador = dialeto.delimiter
    except csv.Error:
        separador = ";" if amostra.count(";") > amostra.count(",") else ","
    leitor = csv.reader(io.StringIO(texto), delimiter=separador)
    linhas = []
    for linha in leitor:
        linhas.append([_texto_celula(v) for v in linha[:MAX_COLUNAS]])
        if len(linhas) > MAX_LINHAS + 200:
            break
    return linhas


def ler_planilha(nome_arquivo: str, dados: bytes) -> Planilha:
    if len(dados) > MAX_BYTES:
        raise PlanilhaInvalida("O arquivo passa de 5 MB. Divida a planilha em partes menores.")
    extensao = nome_arquivo.lower().rsplit(".", 1)[-1] if "." in nome_arquivo else ""
    if extensao == "xlsx":
        if not dados.startswith(b"PK"):
            raise PlanilhaInvalida("O arquivo não é um .xlsx válido.")
        brutas = _ler_xlsx(dados)
    elif extensao == "csv":
        brutas = _ler_csv(dados)
    else:
        raise PlanilhaInvalida("Envie uma planilha .xlsx ou .csv.")

    indice_cabecalho = next((i for i, linha in enumerate(brutas) if any(linha)), None)
    if indice_cabecalho is None:
        raise PlanilhaInvalida("A planilha está vazia.")
    corpo = brutas[indice_cabecalho + 1 :]
    while corpo and not any(corpo[-1]):
        corpo.pop()  # linhas em branco no fim não contam
    cabecalhos = brutas[indice_cabecalho]
    largura = max(len(cabecalhos), *(len(linha) for linha in corpo)) if corpo else len(cabecalhos)
    cabecalhos = [
        c or f"Coluna {i + 1}" for i, c in enumerate(cabecalhos + [""] * (largura - len(cabecalhos)))
    ]
    linhas = [linha + [""] * (largura - len(linha)) for linha in corpo]
    planilha = Planilha(cabecalhos=cabecalhos, linhas=linhas, linha_cabecalho=indice_cabecalho + 1)
    if not planilha.total:
        raise PlanilhaInvalida("A planilha só tem o cabeçalho, sem nenhuma empresa.")
    if planilha.total > MAX_LINHAS:
        raise PlanilhaInvalida(f"A planilha tem {planilha.total} linhas; o limite é {MAX_LINHAS} por lote.")
    return planilha


# --- Mapeamento por regras -----------------------------------------------------


def _sem_acento(texto: str) -> str:
    return unicodedata.normalize("NFKD", texto).encode("ascii", "ignore").decode().lower().strip()


def _parece_site(valor: str) -> bool:
    try:
        normalizar_dominio(valor)
    except EntradaInvalida:
        return False
    return "@" not in valor and " " not in valor.strip()


def mapear_por_regras(planilha: Planilha) -> dict[str, int | None]:
    """Cada coluna vale para no máximo um campo; o conteúdo desempata quando o cabeçalho não diz nada."""
    mapa: dict[str, int | None] = dict.fromkeys(CAMPOS)
    usadas: set[int] = set()
    nomes = [_sem_acento(c) for c in planilha.cabecalhos]
    for campo, padroes in PADROES.items():
        for padrao in padroes:
            achada = next((i for i, n in enumerate(nomes) if i not in usadas and re.search(padrao, n)), None)
            if achada is not None:
                mapa[campo] = achada
                usadas.add(achada)
                break
    if mapa["site"] is None:
        amostra = planilha.linhas[:20]
        for i in range(len(planilha.cabecalhos)):
            valores = [linha[i] for linha in amostra if linha[i]]
            if i not in usadas and valores and sum(_parece_site(v) for v in valores) / len(valores) >= 0.6:
                mapa["site"] = i
                break
    return mapa


def mapeamento_suficiente(mapa: dict[str, int | None]) -> bool:
    return mapa.get("site") is not None or mapa.get("empresa") is not None


def validar_mapeamento(mapa: dict, colunas: int) -> dict[str, int | None]:
    limpo: dict[str, int | None] = {}
    usadas: set[int] = set()
    for campo in CAMPOS:
        valor = mapa.get(campo)
        if valor is None:
            limpo[campo] = None
            continue
        if not isinstance(valor, int) or not 0 <= valor < colunas:
            raise PlanilhaInvalida(f"Coluna inválida para {ROTULO_CAMPO[campo]}.")
        if valor in usadas:
            raise PlanilhaInvalida("A mesma coluna foi escolhida para dois campos.")
        usadas.add(valor)
        limpo[campo] = valor
    if not mapeamento_suficiente(limpo):
        raise PlanilhaInvalida("Escolha pelo menos a coluna da empresa ou a do site.")
    return limpo


# --- Interpretação das linhas -------------------------------------------------


def cnpj_valido(digitos: str) -> bool:
    if len(digitos) != 14 or digitos == digitos[0] * 14:
        return False

    def dv(base: str, pesos: list[int]) -> str:
        resto = sum(int(d) * p for d, p in zip(base, pesos, strict=True)) % 11
        return "0" if resto < 2 else str(11 - resto)

    p1 = [5, 4, 3, 2, 9, 8, 7, 6, 5, 4, 3, 2]
    p2 = [6, *p1]
    return digitos[12] == dv(digitos[:12], p1) and digitos[13] == dv(digitos[:13], p2)


Situacao = Literal["nova", "ja_cadastrada", "repetida", "invalida"]


@dataclass
class LinhaInterpretada:
    numero: int  # número da linha na planilha (o cabeçalho é a linha 1)
    nome: str | None
    dominio: str | None
    cnpj: str | None
    cidade: str | None
    uf: str | None
    contato_nome: str | None
    contato_cargo: str | None
    contato_email: str | None
    situacao: Situacao = "nova"
    motivo: str | None = None
    avisos: list[str] = field(default_factory=list)
    empresa_existente_id: str | None = None
    empresa_existente_nome: str | None = None


def _valor(linha: list[str], mapa: dict, campo: str) -> str:
    i = mapa.get(campo)
    return linha[i].strip() if i is not None and i < len(linha) else ""


def interpretar_linha(numero: int, linha: list[str], mapa: dict) -> LinhaInterpretada:
    avisos: list[str] = []
    site, nome_bruto = _valor(linha, mapa, "site"), _valor(linha, mapa, "empresa")
    dominio = None
    if site:
        try:
            dominio = normalizar_dominio(site)
        except EntradaInvalida:
            avisos.append(f'site "{site[:60]}" ilegível')
    nome = None
    if nome_bruto:
        try:
            nome = normalizar_nome(nome_bruto)
        except EntradaInvalida:
            nome = None
    cnpj_bruto = re.sub(r"\D", "", _valor(linha, mapa, "cnpj"))
    cnpj = None
    if cnpj_bruto:
        cnpj_bruto = cnpj_bruto.zfill(14) if len(cnpj_bruto) in (13, 14) else cnpj_bruto
        if cnpj_valido(cnpj_bruto):
            cnpj = cnpj_bruto
        else:
            avisos.append("CNPJ com dígito verificador errado (ignorado)")
    uf = _valor(linha, mapa, "uf").upper() or None
    if uf and uf not in UFS:
        avisos.append(f'UF "{uf[:10]}" desconhecida (ignorada)')
        uf = None
    email = _valor(linha, mapa, "contato_email").lower() or None
    if email and not EMAIL.match(email):
        avisos.append("e-mail do contato inválido (ignorado)")
        email = None
    r = LinhaInterpretada(
        numero=numero,
        nome=nome,
        dominio=dominio,
        cnpj=cnpj,
        cidade=_valor(linha, mapa, "cidade")[:120] or None,
        uf=uf,
        contato_nome=_valor(linha, mapa, "contato_nome")[:200] or None,
        contato_cargo=_valor(linha, mapa, "contato_cargo")[:200] or None,
        contato_email=email,
        avisos=avisos,
    )
    if not dominio and not nome:
        r.situacao = "invalida"
        r.motivo = "sem nome da empresa nem site legível"
    return r


def numerar(linhas: list[list[str]], linha_cabecalho: int = 1) -> list[tuple[int, list[str]]]:
    """Linhas não vazias com o número que têm no arquivo."""
    return [(linha_cabecalho + 1 + i, linha) for i, linha in enumerate(linhas) if any(linha)]


def classificar(
    linhas: list[list[str]],
    mapa: dict,
    existentes_por_dominio: dict[str, tuple[str, str]],
    existentes_por_cnpj: dict[str, tuple[str, str]],
    linha_cabecalho: int = 1,
) -> list[LinhaInterpretada]:
    """Deduplicação por domínio e CNPJ (§7.2): contra o banco e dentro da própria planilha."""
    vistos_dominio: dict[str, int] = {}
    vistos_cnpj: dict[str, int] = {}
    saida = []
    for i, linha in numerar(linhas, linha_cabecalho):
        r = interpretar_linha(i, linha, mapa)
        if r.situacao == "invalida":
            saida.append(r)
            continue
        existente = (existentes_por_dominio.get(r.dominio) if r.dominio else None) or (
            existentes_por_cnpj.get(r.cnpj) if r.cnpj else None
        )
        repetida = (vistos_dominio.get(r.dominio) if r.dominio else None) or (
            vistos_cnpj.get(r.cnpj) if r.cnpj else None
        )
        if existente:
            r.situacao = "ja_cadastrada"
            r.empresa_existente_id, r.empresa_existente_nome = existente
            r.motivo = f"já cadastrada como {existente[1]}"
        elif repetida:
            r.situacao = "repetida"
            r.motivo = f"repete a linha {repetida}"
        if r.dominio:
            vistos_dominio.setdefault(r.dominio, i)
        if r.cnpj:
            vistos_cnpj.setdefault(r.cnpj, i)
        saida.append(r)
    return saida


def amostra_para_ia(planilha_cabecalhos: list[str], linhas: list[list[str]], n: int = 3) -> list[dict]:
    """O que a IA vê para mapear: cabeçalho e até 3 valores por coluna, com e-mails mascarados."""
    colunas = []
    for i, cab in enumerate(planilha_cabecalhos):
        valores = []
        for linha in linhas[:20]:
            v = linha[i] if i < len(linha) else ""
            if v and len(valores) < n:
                if "@" in v:
                    usuario, _, dominio = v.partition("@")
                    v = f"{usuario[:1]}***@{dominio}"
                valores.append(v[:60])
        colunas.append({"indice": i, "cabecalho": cab, "exemplos": valores})
    return colunas
